# ============================================================================
# ПРОТОТИП (throwaway) — генерация примера xlsx «экспорт всех данных по одному
# объекту недвижимости». Это НЕ продакшн-код: только для показа владельцу
# продукта состава листов и колонок. Не переносить в apps/ без переработки.
#
# Запуск: .venv/bin/python generate_prototype.py
# Результат: PROTO-export-obekt-primer.xlsx (рядом со скриптом)
# ============================================================================
from datetime import date
from pathlib import Path

from openpyxl import Workbook
from openpyxl.styles import Font, PatternFill, Alignment, Border, Side
from openpyxl.utils import get_column_letter

OUT = Path(__file__).parent / "PROTO-export-obekt-primer.xlsx"

PROPERTY_NAME = "2-к квартира, ул. Ленина, 15"

MONEY_FMT = '# ##0.00" ₽"'
DATE_FMT = "DD.MM.YYYY"

HEADER_FONT = Font(bold=True, color="FFFFFF")
HEADER_FILL = PatternFill("solid", fgColor="4472C4")
BLOCK_TITLE_FONT = Font(bold=True, size=12)
THIN = Side(style="thin", color="B0B0B0")
BORDER = Border(left=THIN, right=THIN, top=THIN, bottom=THIN)

# --- Доменные справочники (русские имена из миграции 000078) ----------------
CATEGORIES = {
    "rent": ("Аренда", "income"),
    "utilities": ("Коммунальные услуги", "expense"),
    "repair": ("Ремонт", "expense"),
    "tax": ("Налог", "expense"),
    "deposit_return": ("Возврат депозита", "expense"),
}

LEASE_STATUS_RU = {
    "awaiting_start": "Ожидает начала",
    "active": "Активная",
    "requires_action": "Требует действия",
    "completed": "Завершена",
}

# --- Sample-данные ------------------------------------------------------------
# Арендаторы (TenantContact: Name, Surname, Patronymic, Phone, Email, Comment)
TENANTS = {
    "ivanov": {
        "surname": "Иванов", "name": "Пётр", "patronymic": "Сергеевич",
        "phone": "+79161234567", "email": "p.ivanov@example.ru",
        "comment": "Исторический арендатор, выехал по соглашению сторон",
    },
    "smirnova": {
        "surname": "Смирнова", "name": "Анна", "patronymic": "Владимировна",
        "phone": "+79269876543", "email": "a.smirnova@example.ru",
        "comment": "Просит счета на почту",
    },
    "kozlov": {
        "surname": "Козлов", "name": "Дмитрий", "patronymic": None,
        "phone": "+79031112233", "email": None,
        "comment": None,
    },
}

def fio(t):
    parts = [t["surname"], t["name"], t["patronymic"]]
    return " ".join(p for p in parts if p)

# Аренды (Lease: status, start/end, rent, deposit, payment_day, comment)
LEASES = [
    {"id": "L1", "tenant": TENANTS["ivanov"], "status": "completed",
     "start": date(2025, 2, 1), "end": date(2026, 1, 31),
     "rent": 4500000, "deposit": 4500000, "payment_day": 5,
     "comment": "Договор не продлён"},
    {"id": "L2", "tenant": TENANTS["smirnova"], "status": "active",
     "start": date(2026, 2, 10), "end": date(2027, 2, 9),
     "rent": 4800000, "deposit": 4800000, "payment_day": 10,
     "comment": "С паркингом"},
    {"id": "L3", "tenant": TENANTS["kozlov"], "status": "awaiting_start",
     "start": date(2027, 2, 20), "end": date(2028, 2, 19),
     "rent": 5000000, "deposit": 5000000, "payment_day": 20,
     "comment": ""},
]

LEASE_BY_ID = {l["id"]: l for l in LEASES}

# Операции: только подтверждённые (income->received, expense->paid).
# Поля Operation: type, category(code), status, name, amount(kopecks),
# operation_date, source_operation_date, lease_id(может отсутствовать), comment
OPERATIONS = [
    # февраль 2025 — начало первой аренды
    {"type": "income", "cat": "rent", "status": "received", "name": "Аренда за февраль 2025",
     "amount": 4500000, "date": date(2025, 2, 5), "lease": "L1", "comment": "Получено на карту"},
    {"type": "expense", "cat": "utilities", "status": "paid", "name": "Квитанция ЖКХ февраль 2025",
     "amount": 520000, "date": date(2025, 2, 7), "lease": "L1", "comment": ""},
    {"type": "income", "cat": "rent", "status": "received", "name": "Аренда за март 2025",
     "amount": 4500000, "date": date(2025, 3, 5), "lease": "L1", "comment": ""},
    {"type": "expense", "cat": "utilities", "status": "paid", "name": "Квитанция ЖКХ март 2025",
     "amount": 535000, "date": date(2025, 3, 9), "lease": "L1", "comment": ""},
    {"type": "expense", "cat": "repair", "status": "paid", "name": "Замена смесителя на кухне",
     "amount": 780000, "date": date(2025, 4, 14), "lease": None, "comment": "Операция без привязки к аренде"},
    {"type": "income", "cat": "rent", "status": "received", "name": "Аренда за апрель 2025",
     "amount": 4500000, "date": date(2025, 4, 5), "lease": "L1", "comment": ""},
    {"type": "income", "cat": "rent", "status": "received", "name": "Аренда за май 2025",
     "amount": 4500000, "date": date(2025, 5, 5), "lease": "L1", "comment": ""},
    {"type": "expense", "cat": "tax", "status": "paid", "name": "Налог на имущество за 2024",
     "amount": 960000, "date": date(2025, 6, 1), "lease": None, "comment": ""},
    {"type": "income", "cat": "rent", "status": "received", "name": "Аренда за июнь 2025",
     "amount": 4500000, "date": date(2025, 6, 5), "lease": "L1", "comment": ""},
    {"type": "expense", "cat": "utilities", "status": "paid", "name": "Квитанция ЖКХ июнь 2025",
     "amount": 610000, "date": date(2025, 6, 8), "lease": "L1", "comment": "Летом выше из-за кондиционера"},
    {"type": "income", "cat": "rent", "status": "received", "name": "Аренда за январь 2026",
     "amount": 4500000, "date": date(2026, 1, 5), "lease": "L1", "comment": "Последний платёж Иванова"},
    {"type": "expense", "cat": "deposit_return", "status": "paid", "name": "Возврат депозита Иванову",
     "amount": 4500000, "date": date(2026, 2, 3), "lease": "L1", "comment": "Депозит возвращён полностью"},
    {"type": "income", "cat": "rent", "status": "received", "name": "Аренда за февраль 2026",
     "amount": 4800000, "date": date(2026, 2, 10), "lease": "L2", "comment": ""},
    {"type": "expense", "cat": "utilities", "status": "paid", "name": "Квитанция ЖКХ февраль 2026",
     "amount": 580000, "date": date(2026, 2, 12), "lease": "L2", "comment": ""},
    {"type": "income", "cat": "rent", "status": "received", "name": "Аренда за март 2026",
     "amount": 4800000, "date": date(2026, 3, 10), "lease": "L2", "comment": ""},
]

RU_MONTHS = ["", "Январь", "Февраль", "Март", "Апрель", "Май", "Июнь",
             "Июль", "Август", "Сентябрь", "Октябрь", "Ноябрь", "Декабрь"]

def rub(kopecks):
    return kopecks / 100.0

def style_header_row(ws, row, ncols):
    for c in range(1, ncols + 1):
        cell = ws.cell(row=row, column=c)
        cell.font = HEADER_FONT
        cell.fill = HEADER_FILL
        cell.alignment = Alignment(horizontal="center", vertical="center", wrap_text=True)
        cell.border = BORDER

def set_widths(ws, widths):
    for i, w in enumerate(widths, start=1):
        ws.column_dimensions[get_column_letter(i)].width = w

wb = Workbook()

# ============================ Лист 1: Сводка ==================================
ws = wb.active
ws.title = "Сводка"

total_income = sum(o["amount"] for o in OPERATIONS if o["type"] == "income")
total_expense = sum(o["amount"] for o in OPERATIONS if o["type"] == "expense")

ws["A1"] = f"Объект: {PROPERTY_NAME}"
ws["A1"].font = Font(bold=True, size=14)

ws["A3"] = "Итоги за всё время"
ws["A3"].font = BLOCK_TITLE_FONT
for r, (label, value) in enumerate([
    ("Доход", total_income),
    ("Расход", total_expense),
    ("Прибыль", total_income - total_expense),
], start=4):
    ws.cell(row=r, column=1, value=label).border = BORDER
    c = ws.cell(row=r, column=2, value=rub(value))
    c.number_format = MONEY_FMT
    c.border = BORDER

# --- по месяцам
months = sorted({(o["date"].year, o["date"].month) for o in OPERATIONS})
row = 8
ws.cell(row=row, column=1, value="По месяцам").font = BLOCK_TITLE_FONT
row += 1
for c, h in enumerate(["Месяц", "Доход", "Расход", "Прибыль"], start=1):
    ws.cell(row=row, column=c, value=h)
style_header_row(ws, row, 4)
month_header_row = row
for (y, m) in months:
    row += 1
    inc = sum(o["amount"] for o in OPERATIONS
              if o["type"] == "income" and o["date"].year == y and o["date"].month == m)
    exp = sum(o["amount"] for o in OPERATIONS
              if o["type"] == "expense" and o["date"].year == y and o["date"].month == m)
    ws.cell(row=row, column=1, value=f"{RU_MONTHS[m]} {y}").border = BORDER
    for c, v in enumerate([inc, exp, inc - exp], start=2):
        cell = ws.cell(row=row, column=c, value=rub(v))
        cell.number_format = MONEY_FMT
        cell.border = BORDER

# --- по категориям
row += 2
ws.cell(row=row, column=1, value="По категориям").font = BLOCK_TITLE_FONT
row += 1
for c, h in enumerate(["Категория", "Тип", "Сумма"], start=1):
    ws.cell(row=row, column=c, value=h)
style_header_row(ws, row, 3)
for code in ["rent", "utilities", "repair", "tax", "deposit_return"]:
    ops = [o for o in OPERATIONS if o["cat"] == code]
    if not ops:
        continue
    row += 1
    name, typ = CATEGORIES[code]
    ws.cell(row=row, column=1, value=name).border = BORDER
    ws.cell(row=row, column=2, value="Доход" if typ == "income" else "Расход").border = BORDER
    c = ws.cell(row=row, column=3, value=rub(sum(o["amount"] for o in ops)))
    c.number_format = MONEY_FMT
    c.border = BORDER

set_widths(ws, [32, 16, 16, 16])
ws.freeze_panes = f"A{month_header_row + 1}"  # закрепляем заголовок таблицы по месяцам

# ============================ Лист 2: Операции ================================
ws = wb.create_sheet("Операции")
headers = ["Дата", "Тип", "Категория", "Название", "Сумма", "Аренда / арендатор", "Комментарий"]
for c, h in enumerate(headers, start=1):
    ws.cell(row=1, column=c, value=h)
style_header_row(ws, 1, len(headers))

for r, o in enumerate(OPERATIONS, start=2):
    ws.cell(row=r, column=1, value=o["date"]).number_format = DATE_FMT
    ws.cell(row=r, column=2, value="Доход" if o["type"] == "income" else "Расход")
    ws.cell(row=r, column=3, value=CATEGORIES[o["cat"]][0])
    ws.cell(row=r, column=4, value=o["name"])
    ws.cell(row=r, column=5, value=rub(o["amount"])).number_format = MONEY_FMT
    lease = LEASE_BY_ID.get(o["lease"]) if o["lease"] else None
    ws.cell(row=r, column=6, value=fio(lease["tenant"]) if lease else "—")
    ws.cell(row=r, column=7, value=o["comment"])
    for c in range(1, len(headers) + 1):
        ws.cell(row=r, column=c).border = BORDER

set_widths(ws, [12, 10, 22, 32, 16, 28, 36])
ws.freeze_panes = "A2"
ws.auto_filter.ref = f"A1:G{len(OPERATIONS) + 1}"

# ============================ Лист 3: Аренды ==================================
ws = wb.create_sheet("Аренды")
headers = ["Арендатор", "Статус", "Дата начала", "Дата окончания",
           "Ставка аренды, ₽/мес", "Депозит, ₽", "День платежа", "Комментарий"]
for c, h in enumerate(headers, start=1):
    ws.cell(row=1, column=c, value=h)
style_header_row(ws, 1, len(headers))

for r, l in enumerate(LEASES, start=2):
    ws.cell(row=r, column=1, value=fio(l["tenant"]))
    ws.cell(row=r, column=2, value=LEASE_STATUS_RU[l["status"]])
    ws.cell(row=r, column=3, value=l["start"]).number_format = DATE_FMT
    ws.cell(row=r, column=4, value=l["end"]).number_format = DATE_FMT
    ws.cell(row=r, column=5, value=rub(l["rent"])).number_format = MONEY_FMT
    ws.cell(row=r, column=6, value=rub(l["deposit"])).number_format = MONEY_FMT
    ws.cell(row=r, column=7, value=l["payment_day"])
    ws.cell(row=r, column=8, value=l["comment"])
    for c in range(1, len(headers) + 1):
        ws.cell(row=r, column=c).border = BORDER

set_widths(ws, [28, 16, 14, 14, 18, 16, 13, 28])
ws.freeze_panes = "A2"

# ============================ Лист 4: Арендаторы ==============================
ws = wb.create_sheet("Арендаторы")
headers = ["Фамилия", "Имя", "Отчество", "Телефон", "Email",
           "Связанная аренда / период", "Комментарий"]
for c, h in enumerate(headers, start=1):
    ws.cell(row=1, column=c, value=h)
style_header_row(ws, 1, len(headers))

r = 1
for l in LEASES:
    r += 1
    t = l["tenant"]
    period = (f"{l['start'].strftime('%d.%m.%Y')} — "
              f"{l['end'].strftime('%d.%m.%Y') if l['end'] else 'н/д'} "
              f"({LEASE_STATUS_RU[l['status']]})")
    ws.cell(row=r, column=1, value=t["surname"])
    ws.cell(row=r, column=2, value=t["name"])
    ws.cell(row=r, column=3, value=t["patronymic"] or "")
    ws.cell(row=r, column=4, value=t["phone"] or "")
    ws.cell(row=r, column=5, value=t["email"] or "")
    ws.cell(row=r, column=6, value=period)
    ws.cell(row=r, column=7, value=t["comment"] or "")
    for c in range(1, len(headers) + 1):
        ws.cell(row=r, column=c).border = BORDER

set_widths(ws, [16, 12, 16, 16, 26, 34, 40])
ws.freeze_panes = "A2"

wb.save(OUT)
print(f"OK: {OUT}")
